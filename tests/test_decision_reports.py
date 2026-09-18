import hashlib
import json
import tempfile
from pathlib import Path
import unittest
import zipfile
import xml.etree.ElementTree as ET

from tca_bp.decision_finance import calculate_cap_table
from tca_bp.decision_reports import generate_reports


def sample_snapshot():
    return {"id": "DEMO_FICTIVE", "title": "Projet de services fictif", "as_of": "2026-09-13",
            "profile": {"Activité": "Services de démonstration", "Horizon": "3 années"},
            "metrics": [{"id": "revenue", "label": "Chiffre d’affaires", "value": 120000,
                         "unit": "EUR", "status": "HYPOTHESE", "source_ids": ["S1"]},
                        {"id": "cash", "label": "Trésorerie", "value": 987654321,
                         "unit": "EUR", "status": "NON_VERIFIE", "available": False},
                        {"id": "cost", "label": "Achats", "value": 0, "unit": "EUR", "status": "CONFIRME"}],
            "series": [{"label": "Chiffre d’affaires annuel", "unit": "EUR", "categories": ["2026", "2027", "2028"],
                        "values": [120000, None, 180000], "status": "HYPOTHESE"}],
            "scenarios": [{"name": "Développement", "metrics": [{"label": "Chiffre d’affaires", "value": 150000, "unit": "EUR", "status": "HYPOTHESE"}]}],
            "capital": calculate_cap_table([{"name": "Équipe fictive", "shares": 1000}],
                                           [{"name": "Tour fictif", "pre_money": 4e6, "investment": 1e6}]),
            "sources": [{"id": "S1", "label": "Hypothèses fictives du scénario", "sha256": "a" * 64}],
            "limitations": ["Données entièrement fictives. Aucun résultat de ce rapport ne certifie un dossier réel.",
                            "Les valeurs absentes et les résultats indisponibles restent identifiés."]}


class ReportsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="tca-report-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_ppt_tables_keep_notes_separate_without_losing_rows(self):
        from pptx import Presentation
        from tca_bp.decision_reports import _pptx
        note='Source et qualification conservées dans l’instantané.'
        rows=[[f'Libellé {i} '+'x'*125,str(i)] for i in range(10)]
        path=self.root/'pagination.pptx'
        _pptx(path,'Témoin fictif','Contrôle de pagination',[{'title':'Tableau sourcé','headers':['Libellé','Montant'],'rows':rows,'notes':[note]}],'TCA')
        seen=[]
        for slide in Presentation(path).slides:
            tables=[shape for shape in slide.shapes if shape.has_table]
            if not tables:continue
            notes=[shape for shape in slide.shapes if shape.has_text_frame and shape.text==note]
            self.assertEqual(len(notes),1)
            for table in tables:
                self.assertLess(table.top+table.height,notes[0].top)
                seen.extend([[cell.text for cell in row.cells] for row in list(table.table.rows)[1:]])
        self.assertEqual(seen,rows)

    def test_typed_dates_use_explicit_epoch_not_financial_number_format(self):
        from tca_bp.decision_reports import _sections
        for system,serial,expected in ((1900,46023,'01/01/2026'),(1904,44561,'01/01/2026'),(1900,1,'01/01/1900')):
            snapshot={'excel_date_system':system,'hypotheses':[{'field_id':'start','label':'Début','value':serial,'value_type':'date'}]}
            rows=_sections(snapshot)[0]['rows']
            self.assertEqual(rows[0][1],expected)
        for system,serial in ((None,46023),(1900,0),(1900,60),(1900,-1),(1900,46023.5)):
            rows=_sections({'excel_date_system':system,'hypotheses':[{'field_id':'start','value':serial,'value_type':'date'}]})[0]['rows']
            self.assertIn('Date à vérifier',rows[0][1])

    def test_pdf_long_table_after_full_page_has_no_blank_page_or_orphan_heading(self):
        from pypdf import PdfReader
        from tca_bp.decision_reports import _pdf
        path = self.root / 'long_tables.pdf'
        sections = [
            {'title': 'Flux mensuels', 'headers': ['Mois', 'Valeur'],
             'rows': [[f'Mois_{i:02}', str(i)] for i in range(25)],
             'notes': ['Calculé ; hypothèses qualifiées']},
            {'title': 'Hypothèses détaillées', 'headers': ['Champ', 'Valeur', 'Source'],
             'rows': [[f'Champ_{i:03}', '1000', 'Provenance fictive'] for i in range(80)]},
        ]
        _pdf(path, 'Projet fictif', 'Contrôle de pagination', sections, 'TCA_FOOTER')
        pages = [page.extract_text() or '' for page in PdfReader(path).pages]
        self.assertGreater(len(pages), 2)
        for number, text in enumerate(pages, 1):
            content = [line.strip() for line in text.splitlines()
                       if line.strip() not in ('', 'TCA_FOOTER', str(number))]
            self.assertTrue(content, f'Page {number} vide hors pied de page')
            for title, first_row in (('Flux mensuels', 'Mois_00'),
                                     ('Hypothèses détaillées', 'Champ_000')):
                if title in text:
                    self.assertIn(first_row, text, f'Titre orphelin page {number}')
        full_text = '\n'.join(pages)
        for prefix, count, digits in (('Mois', 25, 2), ('Champ', 80, 3)):
            for index in range(count):
                self.assertEqual(full_text.count(f'{prefix}_{index:0{digits}d}'), 1)

    def test_same_snapshot_all_formats_native_charts_and_tables(self):
        from pypdf import PdfReader
        from docx import Document
        from pptx import Presentation
        snapshot = sample_snapshot()
        before = json.dumps(snapshot, sort_keys=True)
        result = generate_reports(snapshot, self.root)
        self.assertEqual(json.dumps(snapshot, sort_keys=True), before)
        encoded = Path(result["snapshot"]).read_bytes()
        self.assertEqual(hashlib.sha256(encoded).hexdigest(), result["snapshot_sha256"])
        manifest = json.loads(Path(result["manifest"]).read_text(encoding="utf-8"))
        for file in manifest["files"]:
            self.assertEqual(hashlib.sha256((self.root / file["name"]).read_bytes()).hexdigest(), file["sha256"])
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(result["pdf"]).pages)
        document = Document(result["docx"])
        doc_text = "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        presentation = Presentation(result["pptx"])
        ppt_text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
        ppt_text += "\n".join(cell.text for slide in presentation.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells)
        for text in (pdf_text, doc_text, ppt_text):
            self.assertIn("120 000", text)
            self.assertIn("Indisponible", text)
            self.assertNotIn("987 654 321", text)
            self.assertIn("HYPOTHESE", text)
            self.assertIn("S1", text)
        charts = [shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart]
        self.assertEqual(len(charts), 1)
        with zipfile.ZipFile(result["pptx"]) as archive:
            self.assertTrue(any(name.startswith("ppt/embeddings/") and name.endswith(".xlsx") for name in archive.namelist()))
            chart = ET.fromstring(archive.read("ppt/charts/chart1.xml"))
            ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
            points = chart.findall(".//c:val/c:numRef/c:numCache/c:pt", ns)
            self.assertEqual({p.get("idx"): p.find("c:v", ns).text for p in points}, {"0": "120000.0", "2": "180000.0"})
        self.assertFalse(manifest["calculation_performed"])

    def test_no_overwrite_and_invalid_series(self):
        snapshot = sample_snapshot()
        (self.root / "rapport.pdf").write_bytes(b"original")
        with self.assertRaises(FileExistsError):
            generate_reports(snapshot, self.root)
        self.assertEqual((self.root / "rapport.pdf").read_bytes(), b"original")
        snapshot["series"][0]["values"] = [1]
        with self.assertRaises(ValueError):
            generate_reports(snapshot, self.root / "invalid")

    def test_dates_actuals_hypotheses_and_sensitivity(self):
        from docx import Document
        snapshot = sample_snapshot()
        snapshot['metrics'].append({'id':'cash_break_date','label':'Premier mois négatif','value':'2027-04','unit':'mois'})
        snapshot.setdefault('scenarios',[]).append({'name':'Trésorerie positive vérifiée','metrics':[
            {'id':'cash_break_date','label':'Premier mois négatif','value':None,'unit':'mois','status':'CALCULE'}]})
        snapshot['actuals'] = {'cutoff':'2026-01-31','rows':[{'period':'2026-01','metric':'revenue','kind':'flow','value':'9000','unit':'EUR'}],
            'comparisons':[{'period':'2026-01','metric':'revenue','actual':'9000','budget':10000,'variance':-1000}],
            'forecast_series':[{'id':'revenue','label':'Revenus','categories':['2026-01','2026-02'],'values':[9000,12000],'unit':'EUR'}],
            'forecast_status':'A_COMPLETER_OU_RECALCULER','annual_balance':[{'period':'2026','assets':None,'liabilities':None,'equity':None,'net_income':None,'balance_check':None}],
            'diagnostics':[{'code':'SOLDE_MANQUANT','metric':'cash','message':'Trésorerie à compléter'}]}
        snapshot['hypotheses'] = [{'field_id':'x','label':'Prix choisi','value':10,'status':'HYPOTHESE','cell':'A1','evidence_id':'S1'},
                                  {'field_id':'x','label':'Prix choisi','value':20,'status':'HYPOTHESE','cell':'A2','evidence_id':'S1'},
                                  {'field_id':'z','label':'Champ facultatif absent','value':None}]
        snapshot['sensitivities'] = [{'status':'COMPLETE','points':[{'values':[1.2], 'metrics':[{'id':'cash_min','label':'Point bas','value':-5000,'unit':'EUR'}]}]}]
        result=generate_reports(snapshot,self.root)
        document=Document(result['docx'])
        text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
        for expected in ['2027-04','Réalisé validé','9 000','-1 000','Prix choisi','10; 20','Trésorerie à compléter','-5 000',
                         'Bilan et résultat annuels actualisés','sans réimporter les flux déjà inclus','A_COMPLETER_OU_RECALCULER','Aucune sur la période']:
            self.assertIn(expected,text)
        self.assertNotIn('Champ facultatif absent',text)
        max_width=document.sections[0].page_width-document.sections[0].left_margin-document.sections[0].right_margin
        for table in document.tables:
            for row in table.rows:
                self.assertLessEqual(sum(c.width for c in row.cells),max_width+1000)


if __name__ == "__main__":
    unittest.main()
