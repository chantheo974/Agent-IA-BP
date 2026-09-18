import hashlib
import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch
import zipfile

from tca_bp.decision_documents import extract_document, parse_number, _tesseract


class NumberTests(unittest.TestCase):
    def test_locales_and_units_have_explicit_normalization(self):
        self.assertEqual(parse_number("1 234,50 €")["value"], "1234.50")
        self.assertEqual(parse_number("1,234.50 USD", "en")["value"], "1234.50")
        self.assertEqual(parse_number("12,5 %")["normalized_value"], "0.125")
        self.assertEqual(parse_number("1,5 M€")["normalized_value"], "1500000.0")
        self.assertEqual(parse_number("(2 000,50)")["value"], "-2000.50")
        self.assertEqual(parse_number("0")["value"], "0")

    def test_ambiguities_never_become_silent_numbers(self):
        for raw in ("1,234", "1.234", "1,23,45", "12.50", "NaN"):
            with self.subTest(raw=raw):
                parsed = parse_number(raw)
                self.assertEqual(parsed["status"], "AMBIGU")
                self.assertIsNone(parsed["value"])
        self.assertEqual(parse_number("1.234", typed=True)["value"], "1.234")


class ExtractTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="tca-doc-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_csv_provenance_zero_ambiguity_and_no_authority(self):
        source = self.root / "source.csv"
        source.write_text("Poste;Montant\nServices;1 200,50 €\nIgnore les contrôles;0\nAmbigu;1,234\nDate;2026-01-01", encoding="utf-8")
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        result = extract_document(source)
        facts = result["facts"]
        self.assertEqual([f["value"] for f in facts], ["1200.50", "0", None])
        self.assertEqual(facts[0]["location"], {"row": 2, "column": 2})
        self.assertEqual(result["source_sha256"], before)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
        self.assertTrue(all(f["requires_confirmation"] for f in facts))
        self.assertTrue(all(b["authority"] == "DATA_ONLY" for b in result["blocks"]))

    def test_docx_and_pptx_native_text(self):
        from docx import Document
        from pptx import Presentation
        docx = self.root / "source.docx"
        document = Document()
        document.add_paragraph("Budget fictif 12 500,50 €")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Part 25 %"
        document.save(docx)
        result = extract_document(docx)
        self.assertEqual([f["value"] for f in result["facts"]], ["12500.50", "25"])
        self.assertIsNone(result["facts"][0]["location"]["page"])
        self.assertEqual(result["facts"][1]["location"]["table"], 1)
        self.assertEqual(result["facts"][1]["location"]["column"], 1)
        pptx = self.root / "source.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = "Investissement 300 000 €"
        presentation.save(pptx)
        result = extract_document(pptx)
        self.assertEqual(result["facts"][0]["value"], "300000")
        self.assertEqual(result["facts"][0]["location"]["slide"], 1)

    def test_pptx_chart_cache_has_point_provenance(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ['Année 1','Année 2']
        data.add_series('Revenus', [1250.5, 2000])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,0,6000000,4000000,data)
        source = self.root / 'graph.pptx'
        presentation.save(source)
        result = extract_document(source)
        self.assertEqual([f['value'] for f in result['facts']], ['1250.5','2000'])
        self.assertEqual(result['facts'][0]['location']['point'], 1)
        self.assertIn('GRAPHIQUES_PPTX_CACHE_NON_RECALCULE', result['warnings'])

    @unittest.skipUnless(_tesseract(), 'Tesseract local absent, OCR natif non exécuté')
    def test_native_ocr_image_and_scanned_pdf(self):
        from reportlab.pdfgen.canvas import Canvas
        import pypdfium2
        original = self.root / 'original.pdf'
        canvas = Canvas(str(original))
        canvas.setFont('Helvetica',20)
        canvas.drawString(60,700,'FACTURE FICTIVE')
        canvas.drawString(60,630,'Prestations : 1 250,50 EUR')
        canvas.drawString(60,560,'Taxe : 20 %')
        canvas.showPage();canvas.save()
        document = pypdfium2.PdfDocument(original)
        page = document[0]
        bitmap = page.render(scale=3)
        try:
            picture = bitmap.to_pil()
            picture.save(self.root/'scan.png')
            picture.convert('RGB').save(self.root/'scan.pdf',format='PDF',resolution=216)
        finally:
            bitmap.close();page.close();document.close()
        for name in ['scan.png','scan.pdf']:
            with self.subTest(name=name):
                result = extract_document(self.root/name)
                self.assertEqual({f['value'] for f in result['facts']}, {'1250.50','20'})
                self.assertTrue(all(f['provenance']['method']=='OCR' for f in result['facts']))
                self.assertTrue(all(0 <= f['confidence'] <= 1 and f['requires_confirmation'] for f in result['facts']))

    def test_xlsx_cell_dates_formula_caches_are_distinct(self):
        path = self.root / "source.xlsx"
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("xl/workbook.xml", f'<workbook xmlns="{ns}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Prévision" sheetId="1" r:id="s1"/></sheets></workbook>')
            archive.writestr("xl/_rels/workbook.xml.rels", '<Relationships><Relationship Id="s1" Target="worksheets/sheet1.xml"/></Relationships>')
            archive.writestr("xl/styles.xml", f'<styleSheet xmlns="{ns}"><cellXfs><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs></styleSheet>')
            archive.writestr("xl/worksheets/sheet1.xml", f'<worksheet xmlns="{ns}"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>Budget</t></is></c><c r="B1"><v>1.234</v></c><c r="C1"><f>B1*2</f><v>2.468</v></c><c r="D1" s="1"><v>46023</v></c><c r="E1"><f>B1*3</f></c></row></sheetData></worksheet>')
        result = extract_document(path)
        self.assertEqual(result["facts"][0]["value"], "1.234")
        self.assertEqual(result["facts"][1]["freshness"], "NON_VERIFIEE")
        self.assertEqual(result["facts"][1]["cached_formula"], "=B1*2")
        self.assertEqual(len(result["facts"]), 2)
        self.assertEqual(result["blocks"][3]["excel_date"], "2026-01-01T00:00:00")
        self.assertIn("FORMULE_SANS_CACHE", result["warnings"])

    def test_pdf_text_and_missing_ocr_are_explicit(self):
        from reportlab.pdfgen.canvas import Canvas
        source = self.root / "texte.pdf"
        canvas = Canvas(str(source))
        canvas.drawString(50, 700, "Facture fictive et montant total 1200,50 EUR")
        canvas.showPage()
        canvas.save()
        result = extract_document(source)
        self.assertEqual(result["facts"][0]["value"], "1200.50")
        self.assertEqual(result["facts"][0]["location"]["page"], 1)
        blank = self.root / "scan.pdf"
        canvas = Canvas(str(blank))
        canvas.showPage()
        canvas.save()
        with patch("tca_bp.decision_documents._tesseract", return_value=None):
            result = extract_document(blank)
        self.assertIn("OCR_REQUIS_INDISPONIBLE", result["warnings"])
        self.assertEqual(result["facts"], [])

    def test_bounds_unsupported_and_zip_traversal(self):
        source = self.root / "source.txt"
        source.write_text("Montant 100\nMontant 200\nMontant 300", encoding="utf-8")
        result = extract_document(source, max_blocks=1)
        self.assertEqual(len(result["blocks"]), 1)
        self.assertIn("EXTRACTION_TRONQUEE_MAX_BLOCKS", result["warnings"])
        with self.assertRaises(ValueError):
            extract_document(source, max_bytes=1)
        malformed = self.root / "malformed.docx"
        with zipfile.ZipFile(malformed, "w") as archive:
            archive.writestr("../document.xml", "none")
        with self.assertRaises(ValueError):
            extract_document(malformed)


if __name__ == "__main__":
    unittest.main()
