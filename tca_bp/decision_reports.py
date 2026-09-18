"""Three local reports from one immutable JSON snapshot.

The caller supplies qualified metrics and owns the XLSM copy. This module
does not read workbooks, recalculate, access an AI provider, or invent analysis.
PPTX tables and charts are native editable objects with embedded chart data.
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
import datetime as dt
import hashlib
import io
import json
import math
from pathlib import Path
import re
import tempfile
from xml.sax.saxutils import escape


def _json_default(value):
    if isinstance(value, Decimal) and value.is_finite():
        return format(value, "f")
    raise TypeError(f"Valeur non JSON : {type(value).__name__}")


def _text(value, limit=4000):
    if value is None:
        return "À compléter"
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, default=_json_default)
    result = str(value)
    if len(result) > limit or any(ord(c) < 32 and c not in "\n\t\r" for c in result):
        raise ValueError("Texte trop long ou contenant des contrôles interdits")
    return result


def _number(value):
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("Valeur numérique attendue")
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        raise ValueError("Valeur numérique attendue") from None
    if not number.is_finite() or not math.isfinite(float(number)):
        raise ValueError("Valeur numérique finie attendue")
    return number


def _shown(value):
    try:
        number = _number(value)
    except ValueError:
        if isinstance(value, str) and value.strip().lower() not in {"nan", "infinity", "inf", "+inf", "-inf"}:
            return _text(value, 500)
        raise
    if number is None:
        return "À compléter"
    # At most two decimals in public tables; raw precision remains in snapshot.
    text = f"{number:,.2f}".replace(",", " ").replace(".", ",")
    return text[:-3] if text.endswith(",00") else text


def _shown_hash(value):
    text = _text(value)
    return "\n".join(text[i:i + 16] for i in range(0, 64, 16)) if re.fullmatch(r"[0-9a-fA-F]{64}", text) else text


def _hypothesis_shown(question,date_system):
    value=question.get('value')
    if question.get('value_type')!='date' or value is None:
        return _shown(value)
    if isinstance(value,str):
        try:return dt.date.fromisoformat(value).strftime('%d/%m/%Y')
        except ValueError:pass
    try:
        serial=_number(value)
        if date_system not in (1900,1904) or serial!=serial.to_integral_value() or serial<0 or (date_system==1900 and serial in (0,60)):
            raise ValueError('Date ambiguë')
        epoch=dt.date(1904,1,1) if date_system==1904 else dt.date(1899,12,30)
        days=int(serial)+(1 if date_system==1900 and 0<serial<60 else 0)
        return (epoch+dt.timedelta(days=days)).strftime('%d/%m/%Y')
    except (ValueError,OverflowError):
        return 'Date à vérifier ('+_shown(value)+')'


def _metrics(value):
    if isinstance(value, dict):
        return [{"id": key, **item} if isinstance(item, dict) else {"id": key, "label": key, "value": item}
                for key, item in value.items()]
    if not isinstance(value, list):
        raise ValueError("metrics doit être une liste ou un objet")
    return value


def _status(value):
    return _text({'DISPONIBLE_SUR_PREREQUIS_QUALIFIES':'Calculé ; hypothèses qualifiées',
        'DISPONIBLE_SOUS_HYPOTHESES':'Calculé sous hypothèses',
        'A_COMPLETER_OU_RECALCULER':'À compléter ou recalculer','A_PREPARER':'Raccord à préparer',
        'CALCULE':'Calculé','INDISPONIBLE':'Indisponible','NON_VERIFIE':'Non vérifié',
        'EXPLICABLE':'Raccord contrôlé','VARIANTE_METIER_A_REEXAMINER':'Modèle modifié : à qualifier'}.get(value,value))


def _metric_shown(metric):
    if not metric.get('available',True):return 'Indisponible'
    if metric.get('id')=='cash_break_date' and metric.get('value') is None and metric.get('status')=='CALCULE':
        return 'Aucune sur la période'
    return _shown(metric.get('value'))


def _sections(snapshot):
    """One common content model used by all three renderers."""
    sections = []
    profile = snapshot.get("profile", {})
    if profile:
        labels={'company_name':'Entreprise','activity':'Activité','business_model':'Sources de revenus',
                'start_year':'Première année','years':'Horizon (années)','activity_start_month':'Mois de démarrage',
                'goals':'Objectifs','objectives':'Objectifs','modules':'Modules utilisés','description':'Description'}
        sections.append({"title": "Profil du projet", "headers": ["Élément", "Information"],
                         "rows": [[_text(labels.get(k,k)), _text(v)] for k, v in profile.items()]})
    metrics = _metrics(snapshot.get("metrics", []))
    if metrics:
        rows = []
        for m in metrics:
            rows.append([_text(m.get("label", m.get("id", "Indicateur"))),
                         _metric_shown(m),
                         _text(m.get("unit", "")), _status(m.get("status", "NON_VERIFIE")),
                         ", ".join(_text(v) for v in m.get("source_ids", []))])
        sections.append({"title": "Indicateurs du scénario", "headers": ["Indicateur", "Valeur", "Unité", "Qualification", "Sources"], "rows": rows})
    annual=_metrics(snapshot.get('annual_metrics',[]))
    if annual:
        sections.append({'title':'Prévisionnel annuel — horizon complet','headers':['Année','Indicateur','Valeur','Unité','Qualification'],
            'rows':[[_text(m.get('period')),_text(m.get('label',m.get('id'))),_shown(m.get('value')) if m.get('available',True) else 'Indisponible',_text(m.get('unit','EUR')),_status(m.get('status','NON_VERIFIE'))] for m in annual]})
    for series in snapshot.get("series", []):
        categories, values = series.get("categories", []), series.get("values", [])
        if not isinstance(categories, list) or not isinstance(values, list) or len(categories) != len(values) or not 1 <= len(values) <= 240:
            raise ValueError("Chaque série requiert 1 à 240 catégories et autant de valeurs")
        available = series.get("available", True)
        values = [_number(v) if available else None for v in values]
        label, unit = _text(series.get("label", "Évolution")), _text(series.get("unit", ""))
        sections.append({"title": label, "headers": ["Période", f"Valeur {unit}".strip()],
                         "rows": [[_text(k), _shown(v) if available else "Indisponible"] for k, v in zip(categories, values)],
                         "chart": {"label": label, "unit": unit, "categories": [_text(c) for c in categories],
                                   "values": values},
                         "notes": [_status(series.get("status", "NON_VERIFIE"))]})
    scenarios = snapshot.get("scenarios", [])
    for scenario in scenarios:
        rows = [[_text(m.get("label", m.get("id", "Indicateur"))),
                 _metric_shown(m),
                 _text(m.get("unit", "")), _status(m.get("status", "NON_VERIFIE"))]
                for m in _metrics(scenario.get("metrics", []))]
        if rows:
            sections.append({"title": "Scénario " + _text(scenario.get("name", "À compléter")),
                               "headers": ["Indicateur", "Valeur", "Unité", "Qualification"], "rows": rows})
        for section in _sections({'annual_metrics':scenario.get('annual_metrics',[])}):
            if section['title']=='Prévisionnel annuel — horizon complet':
                sections.append({**section,'title':'Scénario '+_text(scenario.get('name','À compléter'))+' — années'})
    capital = snapshot.get("capital") or {}
    final = capital.get("final", capital)
    holders = final.get("shareholders", final.get("holders", []))
    if holders:
        rows = []
        for holder in holders:
            percent = holder.get("percent")
            if percent is None and holder.get("ownership") is not None:
                percent = _number(holder["ownership"]) * 100
            rows.append([_text(holder.get("name", "Détenteur")), _shown(holder.get("shares")), _shown(percent)])
        if "pool_shares" in final:
            rows.append(["Pool réservé", _shown(final["pool_shares"]), _shown(_number(final["pool_ownership"]) * 100)])
        sections.append({"title": "Répartition du capital pleinement dilué", "headers": ["Détenteur", "Parts", "Pourcentage"], "rows": rows,
                         "notes": ["Simulation économique avec parts fractionnaires. Les modalités juridiques restent à confirmer."]})
    if capital.get("rounds"):
        sections.append({"title": "Tours de financement", "headers": ["Tour", "Pré money", "Apport", "Post money", "Pool"],
            "rows": [[_text(r.get("name")), _shown(r.get("pre_money")), _shown(r.get("investment")),
                      _shown(r.get("post_money")), "Avant le tour" if r.get("pool_timing", "before") == "before" else "Après le tour"] for r in capital["rounds"]]})
    actuals = snapshot.get("actuals") or {}
    actual_rows = actuals.get("rows", [])
    metric_labels = {m["id"]: m.get("label", m["id"]) for m in actuals.get("metrics_catalog", [])}
    if actual_rows:
        selected = sorted(actual_rows, key=lambda r: (r.get("period", ""), r.get("metric", "")))[-120:]
        notes = [f"Arrêté au {_text(actuals.get('cutoff'))}. Les flux mensuels et les soldes de clôture conservent des natures distinctes."]
        if len(actual_rows) > len(selected):
            notes.append(f"Les {len(selected)} dernières lignes sont présentées sur {len(actual_rows)}. Le détail complet figure dans l’instantané joint.")
        sections.append({"title": "Réalisé validé", "headers": ["Mois", "Poste", "Nature", "Valeur", "Unité"],
            "rows": [[_text(r.get("period")), _text(metric_labels.get(r.get("metric"), r.get("metric"))),
                      {"flow": "Flux", "balance": "Solde", "volume": "Volume"}.get(r.get("kind"), _text(r.get("kind"))),
                      _shown(r.get("value")), _text(r.get("unit", ""))] for r in selected], "notes": notes})
    comparisons = actuals.get("comparisons", [])
    if comparisons:
        selected = sorted(comparisons, key=lambda r: (r.get("period", ""), r.get("metric", "")))[-120:]
        sections.append({"title": "Écarts par rapport au budget initial", "headers": ["Mois", "Poste", "Budget", "Réalisé", "Écart"],
            "rows": [[_text(r.get("period")), _text(metric_labels.get(r.get("metric"), r.get("metric"))),
                      _shown(r.get("budget")), _shown(r.get("actual")), _shown(r.get("variance"))] for r in selected],
            "notes": ["Écart égal au réalisé moins le budget initial. Les écarts ne constituent pas une analyse causale."] +
                ([f"{len(selected)} dernières comparaisons présentées sur {len(comparisons)} ; détail dans l’instantané."] if len(selected) < len(comparisons) else [])})
    for forecast in actuals.get("forecast_series", []):
        children = _sections({"series": [{**forecast, "label": "Actualisé " + _text(forecast.get("label", forecast.get("id", ""))),
                                          "status": actuals.get("forecast_status", "RACCORD_A_VALIDER")} ]})
        sections.extend(s for s in children if s.get("chart"))
    if actuals.get("forecast_series"):
        sections.append({"title": "Portée de l’actualisé", "paragraphs": [
            "Le passé reprend les données réelles validées. Le futur est lu dans le raccord comptable adopté et recalculé dans Excel : il ajoute aux prévisions du modèle les corrections nettes documentées, sans réimporter les flux déjà inclus.",
            "Trésorerie et postes de raccord sont mensuels ; résultat et bilan futurs sont annuels. Les montants restent indisponibles tant que les sources, décisions de raccord et contrôles ne sont pas satisfaits.",
            "État du raccord : "+_text(actuals.get('forecast_status','A_PREPARER'))]})
    if actuals.get('annual_balance'):
        sections.append({'title':'Bilan et résultat annuels actualisés','headers':['Année','Actif','Passifs hors capitaux propres','Capitaux propres','Résultat net','Écart de bilan'],
            'rows':[[_text(r.get('period')),*[_shown(r.get(k)) for k in ('assets','liabilities','equity','net_income','balance_check')]] for r in actuals['annual_balance']]})
    if actuals.get("diagnostics"):
        sections.append({"title": "Contrôles du raccord au réalisé", "paragraphs": [
            _text(r.get("message") or (str(r.get("code", "À vérifier")) +
                  (" : " + str(r["metric"]) if r.get("metric") else "") +
                  (" ; différence " + str(r["difference"]) if "difference" in r else "") +
                  (" ; mois manquants " + ", ".join(r["periods"]) if r.get("periods") else ""))) for r in actuals["diagnostics"]]})
    hypotheses = snapshot.get("hypotheses", [])
    if isinstance(hypotheses, dict):
        hypotheses = hypotheses.get("questions", [])
    groups = {}
    for question in hypotheses:
        if not isinstance(question, dict):
            continue
        if question.get("value") is None and not question.get("required"):
            continue
        key = (question.get("field_id", question.get("id", question.get("label", ""))), question.get("status", "NON_RENSEIGNE"))
        groups.setdefault(key, []).append(question)
    if groups:
        rows = []
        for (_, status), items in list(groups.items())[:120]:
            first = items[0]
            values = list(dict.fromkeys(_hypothesis_shown(q,snapshot.get('excel_date_system')) for q in items))
            displayed = "; ".join(values[:3])
            if len(values) > 3:
                displayed += f" ; {len(values)} valeurs distinctes au total"
            sources = list(dict.fromkeys(str(q["evidence_id"]) for q in items if q.get("evidence_id")))
            rows.append([_text(first.get("label", first.get("field_id", "Hypothèse"))), displayed,
                         _text(status), str(len(items)), ", ".join(sources[:3]) + ("…" if len(sources) > 3 else "")])
        sections.append({"title": "Hypothèses et informations à compléter", "headers": ["Champ", "Valeur", "Qualification", "Cellules", "Sources"], "rows": rows,
            "notes": ["Les champs renseignés et les informations obligatoires manquantes sont regroupés par champ et qualification. L’instantané conserve chaque cellule."] +
                     ([f"120 groupes présentés sur {len(groups)} ; liste complète dans l’instantané."] if len(groups) > 120 else [])})
    campaigns = snapshot.get("sensitivities", snapshot.get("sensitivity", []))
    if isinstance(campaigns, dict):
        campaigns = [campaigns]
    for campaign in campaigns:
        target_metric = campaign.get("metric", campaign.get("request", {}).get("metric", "cash_min"))
        rows = []
        for index, point in enumerate(campaign.get("points", [])[:100], 1):
            metric = next((m for m in _metrics(point.get("metrics", [])) if m.get("id") == target_metric), None)
            if metric:
                rows.append([str(index), "; ".join(_shown(v) for v in point.get("values", [])), _text(metric.get("label", target_metric)),
                             _shown(metric.get("value")) if metric.get("available", True) else "Indisponible", _text(metric.get("unit", ""))])
        if rows:
            sections.append({"title": "Analyse de sensibilité", "headers": ["Point", "Valeurs des axes", "Indicateur", "Résultat", "Unité"], "rows": rows,
                             "notes": ["Campagne " + _text(campaign.get("status", "NON_VERIFIE")) + ". Les résultats correspondent seulement aux points effectivement calculés."]})
    limitations = snapshot.get("limitations", snapshot.get("limites", []))
    sections.append({"title": "Portée et limites", "paragraphs": [_text(p) for p in limitations] or [
        "Ce rapport restitue l’instantané fourni. Les valeurs absentes restent à compléter et les qualifications indiquées doivent accompagner toute décision."]})
    sources = snapshot.get("sources", [])
    if sources:
        sections.append({"title": "Sources de l’instantané", "headers": ["Référence", "Document", "Empreinte SHA256"],
                         "rows": [[_text(s.get("id", "")), _text(s.get("label", s.get("name", ""))), _shown_hash(s.get("sha256", "À compléter"))] for s in sources]})
    if sum(len(s.get("rows", [])) for s in sections) > 3000 or len(sections) > 100:
        raise ValueError("Rapport trop volumineux pour une génération interactive")
    return sections


def _chart_drawing(chart, start=0, stop=None):
    from reportlab.graphics.shapes import Drawing, String, Line
    from reportlab.graphics.charts.lineplots import LinePlot
    cats = chart["categories"][start:stop]
    vals = chart["values"][start:stop]
    drawing = Drawing(480, 210)
    finite = [(i, float(v)) for i, v in enumerate(vals) if v is not None]
    if not finite:
        drawing.add(String(20, 100, "Valeurs indisponibles", fontName="Helvetica", fontSize=12))
        return drawing
    # Separate segments preserve missing points as gaps instead of zeroes.
    segments, current = [], []
    for i, v in enumerate(vals):
        if v is None:
            if current:
                segments.append(current)
                current = []
        else:
            current.append((i, float(v)))
    if current:
        segments.append(current)
    plot = LinePlot()
    plot.x, plot.y, plot.width, plot.height = 52, 45, 400, 140
    plot.data = segments
    plot.joinedLines = 1
    plot.xValueAxis.valueMin = -0.1
    plot.xValueAxis.valueMax = max(len(cats) - 1, 1) + 0.1
    plot.xValueAxis.valueSteps = list(range(len(cats)))
    plot.xValueAxis.labelTextFormat = lambda n: cats[round(n)] if 0 <= round(n) < len(cats) else ""
    plot.xValueAxis.labels.fontSize = 8
    plot.xValueAxis.labels.angle = 30 if any(len(c) > 8 for c in cats) else 0
    plot.yValueAxis.labels.fontSize = 8
    from reportlab.lib import colors
    from reportlab.graphics.widgets.markers import makeMarker
    for index in range(len(segments)):
        plot.lines[index].strokeColor = colors.HexColor("#244e70")
        plot.lines[index].strokeWidth = 2
        plot.lines[index].symbol = makeMarker("FilledCircle")
        plot.lines[index].symbol.size = 5
    drawing.add(plot)
    drawing.add(String(52, 197, chart["unit"], fontName="Helvetica", fontSize=9))
    return drawing


def _chart_png(chart):
    from reportlab.graphics import renderPDF
    import pypdfium2
    data = renderPDF.drawToString(_chart_drawing(chart))
    document = pypdfium2.PdfDocument(data)
    try:
        page = document[0]
        try:
            bitmap = page.render(scale=2)
            try:
                stream = io.BytesIO()
                bitmap.to_pil().save(stream, format="PNG")
                return stream.getvalue()
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        document.close()


def _pdf(path, title, intro, sections, footer):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    styles = getSampleStyleSheet()
    for name in ("Title", "Heading1", "Heading2"):
        styles[name].textColor = colors.black
    styles['Heading2'].keepWithNext = True
    styles["BodyText"].fontSize, styles["BodyText"].leading = 10, 14
    document = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=1.8 * cm, leftMargin=1.8 * cm,
                                 topMargin=1.8 * cm, bottomMargin=1.8 * cm, title=title, author="TCA BP")
    story = [Paragraph(escape(title), styles["Title"]), Paragraph(escape(intro), styles["BodyText"]), Spacer(1, 15)]
    def para(text):
        return Paragraph(escape(text).replace("\n", "<br/>"), styles["BodyText"])
    for section in sections:
        story.append(Paragraph(escape(section["title"]), styles["Heading2"]))
        if section.get("chart"):
            chart = section["chart"]
            for start in range(0, len(chart["values"]), 12):
                story.append(_chart_drawing(chart, start, start + 12))
        if section.get("rows"):
            headers = section["headers"]
            data = [[para(v) for v in headers]] + [[para(v) for v in row] for row in section["rows"]]
            width = A4[0] - 3.6 * cm
            weights = [2.0] + [1.0] * (len(headers) - 1)
            if section["title"] == "Sources de l’instantané":
                weights = [1, 2, 3]
            table = Table(data, colWidths=[width * w / sum(weights) for w in weights], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7edf2")),
                ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#d9d9d9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            story.append(table)
        for text in section.get("paragraphs", []) + section.get("notes", []):
            story.append(para(text))
        # Heading2 already supplies spaceBefore. A trailing Spacer may spill
        # onto a fresh page and make keepWithNext move a long table once more,
        # leaving that page empty. Keep section spacing on the next heading.
    def page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(1.8 * cm, 1 * cm, footer)
        canvas.drawRightString(A4[0] - 1.8 * cm, 1 * cm, str(doc.page))
        canvas.restoreState()
    document.build(story, onFirstPage=page_number, onLaterPages=page_number)


def _docx(path, title, intro, sections, footer):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    for borders in list(document.styles.element.iter(qn("w:pBdr"))):
        borders.getparent().remove(borders)
    document.styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph(title, "Title")
    document.add_paragraph(intro)
    section.footer.paragraphs[0].text = footer
    section.footer.paragraphs[0].style = document.styles["Caption"]
    for item in sections:
        document.add_heading(item["title"], level=1)
        for note in item.get("notes", []):
            paragraph = document.add_paragraph(note)
            paragraph.paragraph_format.keep_with_next = True
        if item.get("chart"):
            chart = item["chart"]
            for start in range(0, len(chart["values"]), 12):
                chunk = {**chart, "values": chart["values"][start:start + 12], "categories": chart["categories"][start:start + 12]}
                document.add_picture(io.BytesIO(_chart_png(chunk)), width=Inches(6.4))
        if item.get("rows"):
            table = document.add_table(rows=1, cols=len(item["headers"]))
            table.style = "Table Grid"
            table.autofit = False
            weights = [2.0] + [1.0] * (len(item["headers"]) - 1)
            if item["title"] == "Sources de l’instantané":
                weights = [1, 2, 2]
            available_width = section.page_width - section.left_margin - section.right_margin
            # Word stores widths as whole twips (635 EMU); rounding each
            # column upward can make a wide financial table cross the margin.
            widths=[int((int(available_width)//635)*weight/sum(weights))*635 for weight in weights]
            for column, width in zip(table.columns, widths):
                column.width = width
            for cell, text in zip(table.rows[0].cells, item["headers"]):
                cell.text = text
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E7EDF2")
                cell._tc.get_or_add_tcPr().append(shading)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in item["rows"]:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for row in table.rows:
                for column_index, cell in enumerate(row.cells):
                    cell.width = widths[column_index]
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(5)
                        paragraph.paragraph_format.space_before = Pt(5)
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                border = OxmlElement(f"w:{edge}")
                for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    border.set(qn("w:" + key), val)
                borders.append(border)
            table._tbl.tblPr.append(borders)
        for text in item.get("paragraphs", []):
            document.add_paragraph(text)
    document.core_properties.title = title
    document.core_properties.author = "TCA BP"
    document.save(path)


def _pptx(path, title, intro, sections, footer):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    blank = presentation.slide_layouts[6]
    def text(slide, content, x, y, w, h, size=18, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.text = content
        for paragraph in frame.paragraphs:
            paragraph.font.name = "Calibri"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor(0, 0, 0)
        return shape
    def slide_for(heading):
        slide = presentation.slides.add_slide(blank)
        text(slide, heading, .65, .4, 12, 1.1, 30, True)
        text(slide, footer, .65, 7.06, 12, .22, 9)
        return slide
    cover = slide_for(title)
    text(cover, intro, .75, 2.3, 11.5, 2.5, 23)
    for item in sections:
        if item.get("chart"):
            chart = item["chart"]
            for start in range(0, len(chart["values"]), 12):
                slide = slide_for(item["title"])
                values, categories = chart["values"][start:start + 12], chart["categories"][start:start + 12]
                if any(v is not None for v in values):
                    data = CategoryChartData()
                    data.categories = categories
                    data.add_series(chart["label"], tuple(None if v is None else float(v) for v in values))
                    graph = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.8), Inches(1.6),
                                                   Inches(11.7), Inches(4.75), data).chart
                    graph.has_legend = False
                    graph.has_title = False
                    graph.category_axis.tick_labels.font.size = Pt(14)
                    graph.value_axis.tick_labels.font.size = Pt(14)
                    graph.chart_style = 13
                else:
                    text(slide, "Valeurs indisponibles", 1, 3, 11, 1, 24)
                text(slide, chart["unit"] + "   " + "; ".join(item.get("notes", [])), .8, 6.45, 11.5, .45, 16)
        rows = item.get("rows", [])
        if rows:
            headers = item["headers"]
            # Row budget follows text length, preserving long evidence labels.
            chunks, chunk, height = [], [], 0.0
            body_height = 4.2 if item.get("notes") else 4.7
            for row in rows:
                max_lines = max(sum(max(1, math.ceil(len(line) / (62 if len(headers) == 2 else 35)))
                                    for line in value.split("\n")) for value in row)
                row_height = max(.48, .27 * max_lines + .12)
                if row_height > body_height:
                    raise ValueError("Une cellule est trop longue pour un tableau PowerPoint lisible")
                if chunk and height + row_height > body_height:
                    chunks.append(chunk)
                    chunk, height = [], 0
                chunk.append((row, row_height))
                height += row_height
            if chunk:
                chunks.append(chunk)
            for index, chunk in enumerate(chunks, 1):
                heading = item["title"] + (f"  {index}/{len(chunks)}" if len(chunks) > 1 else "")
                slide = slide_for(heading)
                table = slide.shapes.add_table(len(chunk) + 1, len(headers), Inches(.65), Inches(1.55),
                                               Inches(12), Inches(.55 + sum(h for _, h in chunk))).table
                table.rows[0].height = Inches(.55)
                for row_index, row in enumerate([headers] + [r for r, _ in chunk]):
                    if row_index:
                        table.rows[row_index].height = Inches(chunk[row_index - 1][1])
                    for col, value in enumerate(row):
                        cell = table.cell(row_index, col)
                        cell.text = value
                        cell.margin_left = cell.margin_right = Inches(.10)
                        cell.margin_top = cell.margin_bottom = Inches(.06)
                        if row_index == 0:
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = RGBColor.from_string("E7EDF2")
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.name, paragraph.font.size = "Calibri", Pt(17)
                            paragraph.font.bold = row_index == 0
                            paragraph.font.color.rgb = RGBColor(0, 0, 0)
                slide.notes_slide.notes_text_frame.text = "\n".join(item.get("notes", [])) + "\n" + footer
                if item.get("notes"):
                    text(slide, " ".join(item["notes"]), .8, 6.45, 11.7, .52, 15)
        paragraphs = item.get("paragraphs", [])
        for offset in range(0, len(paragraphs), 5):
            slide = slide_for(item["title"])
            page_text = "\n\n".join(paragraphs[offset:offset + 5])
            if len(page_text) > 1800:
                raise ValueError("Paragraphe trop long pour une diapositive lisible")
            text(slide, page_text, .8, 1.6, 11.7, 5.1, 19)
    presentation.core_properties.title = title
    presentation.core_properties.author = "TCA BP"
    presentation.save(path)


def generate_reports(snapshot: dict, output_dir) -> dict:
    """Create PDF/DOCX/PPTX plus the exact snapshot and a SHA256 manifest.

Snapshot keys: id, title, as_of, profile (object), metrics (list or mapping),
series [{label,unit,categories,values,status,available}], scenarios
[{name,metrics}], capital (calculate_cap_table result), sources
[{id,label,sha256}], limitations [text]. A metric has id/label/value/unit/
status/source_ids/available. Null stays absent; available=False masks values.
Generation refuses existing output filenames, preventing silent replacement.
"""
    if not isinstance(snapshot, dict):
        raise ValueError("Un instantané objet est requis")
    encoded = json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
                         allow_nan=False, default=_json_default).encode("utf-8")
    if len(encoded) > 5 * 1024 * 1024:
        raise ValueError("Instantané supérieur à 5 Mio")
    frozen = json.loads(encoded)
    digest = hashlib.sha256(encoded).hexdigest()
    title = _text(frozen.get("title", "Rapport de décision financière"), 180)
    identifier = _text(frozen.get("id", digest[:12]), 200)
    date = _text(frozen.get("as_of", "date à compléter"), 100)
    intro = f"Ce rapport présente le projet et ses scénarios à partir de l’instantané {identifier}, arrêté au {date}. Les tableaux conservent les qualifications et les limites des informations fournies."
    sections = _sections(frozen)
    footer = f"TCA BP  |  {digest[:16]}  |  {date}"
    destination = Path(output_dir).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    names = {"pdf": "rapport.pdf", "docx": "rapport.docx", "pptx": "presentation.pptx",
             "snapshot": "instantane.json", "manifest": "manifest.json"}
    if any((destination / filename).exists() for filename in names.values()):
        raise FileExistsError("La destination contient déjà un livrable de ce nom")
    # Author in an isolated temporary directory; failed generation leaves no
    # apparent successful delivery. Each publication uses exclusive creation.
    with tempfile.TemporaryDirectory(prefix="tca-reports-", dir=destination) as scratch:
        temporary = Path(scratch)
        _pdf(temporary / names["pdf"], title, intro, sections, footer)
        _docx(temporary / names["docx"], title, intro, sections, footer)
        _pptx(temporary / names["pptx"], title, intro, sections, footer)
        (temporary / names["snapshot"]).write_bytes(encoded)
        files = [{"format": key, "name": name, "sha256": hashlib.sha256((temporary / name).read_bytes()).hexdigest(),
                  "bytes": (temporary / name).stat().st_size} for key, name in names.items() if key != "manifest"]
        manifest = {"schema_version": "tca-reports/1", "snapshot_sha256": digest, "snapshot_id": identifier,
                    "files": files, "calculation_performed": False, "pptx_native_tables_and_charts": True,
                    "financial_validation": "QUALIFICATIONS_DE_L_INSTANTANE_FOURNI"}
        (temporary / names["manifest"]).write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        published = []
        try:
            for name in names.values():
                target = destination / name
                with target.open("xb") as stream:
                    published.append(target)
                    stream.write((temporary / name).read_bytes())
        except Exception:
            for target in published:
                target.unlink(missing_ok=True)
            raise
    return {**{key: str(destination / name) for key, name in names.items()}, "snapshot_sha256": digest}
